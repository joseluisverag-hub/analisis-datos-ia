from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any, Dict, List

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


def load_json(path: Path) -> Dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(f"No existe: {path}")
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def safe_text(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def split_long_text(text: str, max_len: int = 180) -> List[str]:
    clean = safe_text(text)
    if not clean:
        return []
    chunks: List[str] = []
    remaining = clean
    while len(remaining) > max_len:
        cut = remaining.rfind(" ", 0, max_len)
        if cut <= 0:
            cut = max_len
        chunks.append(remaining[:cut].strip())
        remaining = remaining[cut:].strip()
    if remaining:
        chunks.append(remaining)
    return chunks


def pick_conclusion(manifest: Dict[str, Any]) -> str:
    conclusion = safe_text(manifest.get("conclusion") or "")
    if conclusion:
        return conclusion
    return "No se encontró una conclusión explícita en el notebook."


def add_title_slide(prs: Presentation, manifest: Dict[str, Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    colors = {"navy": "183B56", "gray": "6B7280", "dark": "1F2937"}

    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0), Inches(13.333), Inches(0.5)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor.from_string(colors["navy"])
    header.line.color.rgb = RGBColor.from_string(colors["navy"])

    tx = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(11.5), Inches(0.8))
    p = tx.text_frame.paragraphs[0]
    p.text = safe_text(manifest.get("title") or "Reporte ejecutivo")
    p.font.name = "Aptos Display"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(colors["navy"])

    st = slide.shapes.add_textbox(Inches(0.7), Inches(1.9), Inches(11.5), Inches(0.4))
    sp = st.text_frame.paragraphs[0]
    sp.text = "Documento generado automáticamente desde la extracción del notebook"
    sp.font.name = "Aptos"
    sp.font.size = Pt(13)
    sp.font.color.rgb = RGBColor.from_string(colors["gray"])

    nb = slide.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(11.5), Inches(0.4))
    np = nb.text_frame.paragraphs[0]
    np.text = f"Notebook origen: {safe_text(manifest.get('notebook_path') or 'N/D')}"
    np.font.name = "Aptos"
    np.font.size = Pt(11)
    np.font.color.rgb = RGBColor.from_string(colors["dark"])

    info = slide.shapes.add_textbox(Inches(0.7), Inches(2.9), Inches(11.5), Inches(0.4))
    ip = info.text_frame.paragraphs[0]
    insights = manifest.get("insights") or []
    charts = manifest.get("charts") or []
    ip.text = f"Insights detectados: {len(insights)} | Gráficos detectados: {len(charts)}"
    ip.font.name = "Aptos"
    ip.font.size = Pt(11)
    ip.font.color.rgb = RGBColor.from_string(colors["dark"])

    body = slide.shapes.add_textbox(Inches(0.7), Inches(3.5), Inches(11.5), Inches(0.6))
    bp = body.text_frame.paragraphs[0]
    bp.text = "Estructura: título, insights, conclusión y gráficos extraídos automáticamente."
    bp.font.name = "Aptos"
    bp.font.size = Pt(15)
    bp.font.color.rgb = RGBColor.from_string(colors["dark"])


def add_insights_slide(prs: Presentation, insights: List[str]) -> None:
    navy = RGBColor.from_string("183B56")
    dark = RGBColor.from_string("1F2937")

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12), Inches(0.5))
    p = title.text_frame.paragraphs[0]
    p.text = "Insights clave"
    p.font.name = "Aptos Display"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = navy

    y = 1.2
    for idx, ins in enumerate(insights[:6], 1):
        box = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.8), Inches(0.55))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"{idx}. {safe_text(ins)}"
        p.font.name = "Aptos"
        p.font.size = Pt(13)
        p.font.color.rgb = dark
        y += 0.75


def add_chart_slide(prs: Presentation, chart_path: Path, caption: str, insight: str, source_cell: Any) -> None:
    navy = RGBColor.from_string("183B56")
    dark = RGBColor.from_string("1F2937")
    gray = RGBColor.from_string("6B7280")
    light = RGBColor.from_string("F6F8FB")
    border = RGBColor.from_string("D8E2EE")

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(12), Inches(0.5))
    p = title.text_frame.paragraphs[0]
    p.text = safe_text(caption or "Gráfico")
    p.font.name = "Aptos Display"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = navy

    slide.shapes.add_picture(str(chart_path), Inches(0.7), Inches(1.2), width=Inches(7.0), height=Inches(4.6))

    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.3), Inches(4.7), Inches(4.4))
    panel.fill.solid()
    panel.fill.fore_color.rgb = light
    panel.line.color.rgb = border

    t = slide.shapes.add_textbox(Inches(8.25), Inches(1.55), Inches(3.8), Inches(0.3))
    pt = t.text_frame.paragraphs[0]
    pt.text = "Insight asociado"
    pt.font.name = "Aptos"
    pt.font.size = Pt(14)
    pt.font.bold = True
    pt.font.color.rgb = navy

    chunks = split_long_text(insight, 220)
    if not chunks:
        chunks = ["No se encontró insight asociado en la extracción."]

    iy = 2.0
    for chunk in chunks[:5]:
        ib = slide.shapes.add_textbox(Inches(8.25), Inches(iy), Inches(4.0), Inches(0.55))
        ip = ib.text_frame.paragraphs[0]
        ip.text = f"• {chunk}"
        ip.font.name = "Aptos"
        ip.font.size = Pt(12)
        ip.font.color.rgb = dark
        iy += 0.7

    src = slide.shapes.add_textbox(Inches(8.25), Inches(5.2), Inches(3.5), Inches(0.25))
    sp = src.text_frame.paragraphs[0]
    sp.text = f"Fuente: celda {source_cell}"
    sp.font.name = "Aptos"
    sp.font.size = Pt(10)
    sp.font.color.rgb = gray


def add_conclusion_slide(prs: Presentation, conclusion: str) -> None:
    navy = RGBColor.from_string("183B56")
    dark = RGBColor.from_string("1F2937")
    bg = RGBColor.from_string("FBFCFE")
    border = RGBColor.from_string("D8E2EE")

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title = slide.shapes.add_textbox(Inches(0.6), Inches(0.5), Inches(12), Inches(0.5))
    p = title.text_frame.paragraphs[0]
    p.text = "Conclusión"
    p.font.name = "Aptos Display"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = navy

    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.75), Inches(1.3), Inches(11.7), Inches(4.7))
    box.fill.solid()
    box.fill.fore_color.rgb = bg
    box.line.color.rgb = border

    body = slide.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.0), Inches(4.0))
    tf = body.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = safe_text(conclusion)
    p.font.name = "Aptos"
    p.font.size = Pt(16)
    p.font.color.rgb = dark


def build_pptx(manifest_path: str, charts_dir: str, output_pptx: str) -> None:
    manifest = load_json(Path(manifest_path))
    charts_base = Path(charts_dir)
    if not charts_base.exists():
        raise FileNotFoundError(f"No existe la carpeta de gráficos: {charts_base}")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.author = "OpenAI"
    prs.core_properties.title = safe_text(manifest.get("title") or "Reporte ejecutivo")
    prs.core_properties.subject = "Reporte ejecutivo generado desde notebook"

    insights = manifest.get("insights") or []
    charts = manifest.get("charts") or []
    conclusion = pick_conclusion(manifest)

    add_title_slide(prs, manifest)
    add_insights_slide(prs, insights)

    for idx, chart in enumerate(charts):
        chart_file = charts_base / chart["filename"]
        if not chart_file.exists():
            continue
        caption = safe_text(chart.get("caption") or f"Gráfico {idx + 1}")
        insight = safe_text(insights[idx] if idx < len(insights) else "")
        add_chart_slide(prs, chart_file, caption, insight, chart.get("source_cell_index", "N/D"))

    add_conclusion_slide(prs, conclusion)
    prs.save(output_pptx)


def main() -> None:
    if len(sys.argv) < 4:
        print("Uso: python generate_pptx_from_framework.py <manifest.json> <charts_dir> <output.pptx>")
        sys.exit(1)

    manifest_path = sys.argv[1]
    charts_dir = sys.argv[2]
    output_pptx = sys.argv[3]

    build_pptx(manifest_path, charts_dir, output_pptx)
    print(f"PPTX generado: {output_pptx}")


if __name__ == "__main__":
    main()
